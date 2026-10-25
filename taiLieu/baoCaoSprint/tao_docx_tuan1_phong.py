import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_inline_formatted_text(paragraph, text, base_font_size=12, default_color=RGBColor(0x22, 0x22, 0x22)):
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = default_color
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = default_color
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(0x8F, 0x44, 0x0F)
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = default_color

def create_individual_report():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(base_dir, "baoCaoTuan1_Phong.md")
    docx_path = os.path.join(base_dir, "baoCaoTuan1_Phong.docx")

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Standard margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.79)
        sec.bottom_margin = Inches(0.79)
        sec.left_margin = Inches(1.18)
        sec.right_margin = Inches(0.79)

    # Base Normal style
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(4)

    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        # Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            parsed_rows = []
            for tline in table_lines:
                cells = [c.strip() for c in tline.strip('|').split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                parsed_rows.append(cells)

            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="B0B0B0")

                for r_idx, row_data in enumerate(parsed_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell = tbl.cell(r_idx, c_idx)
                        text = row_data[c_idx] if c_idx < len(row_data) else ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.15

                        clean_text = text.replace('**', '')
                        run = p.add_run(clean_text)
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(10.5)

                        if is_header:
                            set_cell_background(cell, "8F440F") # Màu nâu mộc Quán Mộc
                            run.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "FDFBF8")
                            else:
                                set_cell_background(cell, "FFFFFF")
                            
                            if "Hoàn thành" in text:
                                run.bold = True
                                run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
                            else:
                                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_after = Pt(4)
            continue

        # Main Title #
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x8F, 0x44, 0x0F)
            i += 1
            continue

        # Subtitle **...**
        if stripped.startswith('**HỌC PHẦN:') or stripped.startswith('**DỰ ÁN:'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            clean_text = stripped.replace('**', '')
            run = p.add_run(clean_text)
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x2C, 0x18, 0x10)
            i += 1
            continue

        # Heading 3 (Roman numerals ### I., ### II.)
        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x8F, 0x44, 0x0F)
            i += 1
            continue

        # Heading 4 (#### 1., #### 2.)
        if stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(stripped[5:])
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x2C, 0x18, 0x10)
            i += 1
            continue

        # Bullet lists (- ...)
        if stripped.startswith('- '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            run_bullet = p.add_run('• ')
            run_bullet.bold = True
            run_bullet.font.name = 'Times New Roman'
            add_inline_formatted_text(p, stripped[2:])
            i += 1
            continue

        # Numbered lists (1. , 2. )
        m_num = re.match(r'^(\d+\.)\s+(.*)$', stripped)
        if m_num:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            run_num = p.add_run(m_num.group(1) + ' ')
            run_num.bold = True
            run_num.font.name = 'Times New Roman'
            add_inline_formatted_text(p, m_num.group(2))
            i += 1
            continue

        # Regular paragraphs
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        if "Hà Nội, ngày" in stripped or "Người làm báo cáo" in stripped or "(Ký và ghi rõ họ tên)" in stripped or stripped == "Phong":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_inline_formatted_text(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"Đã tạo thành công tài liệu Word: {docx_path}")

if __name__ == '__main__':
    create_individual_report()
